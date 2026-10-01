import openpyxl

def bubble_sort(data, column):
    n = len(data)
    for i in range(n):
        for j in range(0, n - i - 1):
            if data[j][column] > data[j+1][column]:
                data[j], data[j+1] = data[j+1], data[j]
    return data

# Quy trình xử lý file excel
# B1: Đọc dữ liệu Excel
workbook = openpyxl.load_workbook("visitor_data.xlsx")
worksheet = workbook["Sheet1"]

# B2: Đưa các hàng dữ liệu vào Python
data = []

for row in worksheet.iter_rows(
    min_row = 2,
    values_only = True
):
    data.append(list(row))

# B3: Sắp xếp dữ liệu theo Pricing Plan ID
sorted_data = bubble_sort(data, 6)

# B4: Ghi dữ liệu đã sắp xếp trở lại worksheet
for row_index, row in enumerate(sorted_data, start=2):
    for column_index, value in enumerate(row):
        worksheet.cell(
            row = row_index,
            column=column_index + 1,
            value=value
        )
# B5: Lưu thành file mới 
workbook.save("visitor.xlsx")